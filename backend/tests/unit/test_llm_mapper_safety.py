from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from assembly.column_semantics import (
    ColumnSemanticMap,
    _deterministic_column_roles,
    build_column_semantic_map,
)
from assembly.llm_mapper import (
    apply_cell_mapping,
    sanitize_cell_mapping,
    template_to_structure_text,
)
from assembly.llm_mapper import _plan_call_chunks
from schemas.models import ISOStandard


def _clause_table_document() -> tuple[Document, ColumnSemanticMap]:
    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    table.cell(0, 0).text = "ISO 9001:2015"
    table.cell(1, 0).text = "Requirements"
    table.cell(1, 1).text = "Findings"
    table.cell(1, 2).text = "Conclusion (✓ / NC / OBS)"
    table.cell(2, 0).text = "8.7 Control of nonconforming outputs"
    table.cell(3, 0).text = "9.1 Monitoring, measurement, analysis and evaluation"

    # Blank template cells can still carry the form's intended paragraph/run style.
    finding_paragraph = table.cell(2, 1).paragraphs[0]
    finding_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    finding_paragraph.add_run("").bold = True

    semantic_map = ColumnSemanticMap(
        table_col_roles={
            1: {1: "clause_ref", 2: "findings", 3: "conclusion"},
        }
    )
    return doc, semantic_map


def test_wrong_clause_coordinate_is_relocated_to_findings_and_label_is_preserved():
    doc, semantic_map = _clause_table_document()
    narrative = (
        "The audit team reviewed PR-014 and FR-015 and confirmed that monitoring "
        "controls were implemented and retained as documented information."
    )
    mapping = {"T1_R3_C1": narrative}

    filled = apply_cell_mapping(doc.element.body, mapping, semantic_map=semantic_map)

    table = doc.tables[0]
    assert filled == 2  # repaired finding plus auto-added conclusion tick
    assert table.cell(2, 0).text == "8.7 Control of nonconforming outputs"
    assert table.cell(2, 1).text == narrative
    assert table.cell(2, 2).text == "√"
    assert "T1_R3_C1" not in mapping
    assert mapping["T1_R3_C2"] == narrative


def test_narrative_in_conclusion_column_is_relocated_to_findings():
    doc, semantic_map = _clause_table_document()
    narrative = "Objective evidence was sampled from the internal audit programme and records."

    safe = sanitize_cell_mapping(
        doc.element.body,
        {"T1_R4_C3": narrative},
        semantic_map=semantic_map,
    )

    assert safe == {"T1_R4_C2": narrative}


def test_conclusion_words_are_normalized_to_template_tokens():
    doc, semantic_map = _clause_table_document()

    safe = sanitize_cell_mapping(
        doc.element.body,
        {"T1_R3_C3": "Conforming", "T1_R4_C3": "Observation"},
        semantic_map=semantic_map,
    )

    assert safe["T1_R3_C3"] == "√"
    assert safe["T1_R4_C3"] == "OBS"


def test_direct_findings_mapping_wins_over_relocated_duplicate():
    doc, semantic_map = _clause_table_document()
    wrong = "This narrative was proposed for the pre-printed clause cell and must not win."
    correct = "Correct finding mapped directly to the Findings column."

    safe = sanitize_cell_mapping(
        doc.element.body,
        {"T1_R3_C1": wrong, "T1_R3_C2": correct},
        semantic_map=semantic_map,
    )

    assert safe == {"T1_R3_C2": correct}


def test_fill_preserves_existing_blank_cell_paragraph_and_run_formatting():
    doc, semantic_map = _clause_table_document()

    apply_cell_mapping(
        doc.element.body,
        {"T1_R3_C2": "Styled finding."},
        semantic_map=semantic_map,
    )

    paragraph = doc.tables[0].cell(2, 1).paragraphs[0]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.runs[0].bold is True


def test_deterministic_roles_scan_late_tables_beyond_old_excerpt_limit():
    prefix = "X" * 7000
    structure = "\n".join(
        [
            prefix,
            "TABLE 19",
            "  T19_R2_C1: Requirements [LABEL — DO NOT MODIFY]",
            "  T19_R2_C2: Findings [LABEL — DO NOT MODIFY]",
            "  T19_R2_C3: Conclusion (✓ / NC / OBS) [LABEL — DO NOT MODIFY]",
        ]
    )

    roles = _deterministic_column_roles(structure)

    assert roles[19] == {1: "clause_ref", 2: "findings", 3: "conclusion"}


def test_explicit_template_headers_override_incorrect_ai_classification():
    structure = "\n".join(
        [
            "TABLE 19",
            "  T19_R2_C1: Requirements [LABEL — DO NOT MODIFY]",
            "  T19_R2_C2: Findings [LABEL — DO NOT MODIFY]",
            "  T19_R2_C3: Conclusion (✓ / NC / OBS) [LABEL — DO NOT MODIFY]",
        ]
    )

    class FakeMessages:
        def create(self, **_kwargs):
            return SimpleNamespace(
                content=[
                    SimpleNamespace(
                        text='{"tables":{"19":{"1":"findings","2":"conclusion","3":"clause_ref"}}}'
                    )
                ]
            )

    client = SimpleNamespace(messages=FakeMessages())
    result = build_column_semantic_map(structure, client, "test-model")

    assert result.table_col_roles[19] == {
        1: "clause_ref",
        2: "findings",
        3: "conclusion",
    }


def test_split_table_chunks_repeat_actual_column_headers(tmp_path):
    doc = Document()
    table = doc.add_table(rows=75, cols=3)
    table.cell(0, 0).text = "ISO 9001:2015"
    table.cell(1, 0).text = "Requirements"
    table.cell(1, 1).text = "Findings"
    table.cell(1, 2).text = "Conclusion (✓ / NC / OBS)"
    for row_index in range(2, 75):
        table.cell(row_index, 0).text = f"Clause {row_index} requirement"
    template = tmp_path / "large-template.docx"
    doc.save(template)

    chunks = _plan_call_chunks(str(template), [ISOStandard.QMS])
    continuation = next(chunk for chunk in chunks if "rows36-70" in chunk["label"])

    assert "T1_R1_C1: ISO 9001:2015" in continuation["structure_text"]
    assert "T1_R2_C1: Requirements" in continuation["structure_text"]
    assert "T1_R2_C2: Findings" in continuation["structure_text"]
    assert "T1_R2_C3: Conclusion" in continuation["structure_text"]


def test_real_fr232_qms_clause_narrative_cannot_overwrite_requirements_column():
    template = (
        Path(__file__).resolve().parents[2]
        / "uaf_blank_set"
        / "9-14-45-22-5001"
        / "Surveillance"
        / "FR.232_Audit_Report_R12&09.10.2025.docx"
    )
    doc = Document(template)
    structure = template_to_structure_text(str(template), [ISOStandard.QMS])
    semantic_map = ColumnSemanticMap(
        table_col_roles=_deterministic_column_roles(structure)
    )
    original_requirement = doc.tables[18].cell(47, 0).text
    narrative = (
        "The quality-control procedure and sampled records were reviewed; the audit "
        "team confirmed that nonconforming outputs were identified and controlled."
    )

    mapping = {"T19_R48_C1": narrative}
    apply_cell_mapping(doc.element.body, mapping, semantic_map=semantic_map)

    assert doc.tables[18].cell(47, 0).text == original_requirement
    assert doc.tables[18].cell(47, 1).text == narrative
    assert doc.tables[18].cell(47, 2).text == "√"


def test_real_fr232_compound_clause_block_is_split_into_matching_rows():
    template = (
        Path(__file__).resolve().parents[2]
        / "uaf_blank_set"
        / "9-14-45-22-5001"
        / "Surveillance"
        / "FR.232_Audit_Report_R12&09.10.2025.docx"
    )
    doc = Document(template)
    structure = template_to_structure_text(str(template), [ISOStandard.QMS])
    semantic_map = ColumnSemanticMap(
        table_col_roles=_deterministic_column_roles(structure)
    )
    compound = """8.7 — Conforming (√). Nonconforming outputs and sampled records were controlled.

9.1 — Conforming (√). Monitoring objectives and product results were evaluated.

9.2 — Observation. Internal audit arrangements were implemented at planned intervals."""

    mapping = {"T19_R48_C1": compound}
    apply_cell_mapping(doc.element.body, mapping, semantic_map=semantic_map)

    qms_table = doc.tables[18]
    assert qms_table.cell(47, 1).text == "Nonconforming outputs and sampled records were controlled."
    assert qms_table.cell(48, 1).text == "Monitoring objectives and product results were evaluated."
    assert qms_table.cell(51, 1).text == "Internal audit arrangements were implemented at planned intervals."
    assert qms_table.cell(47, 2).text == "√"
    assert qms_table.cell(48, 2).text == "√"
    assert qms_table.cell(51, 2).text == "OBS"
